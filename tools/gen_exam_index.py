#!/usr/bin/env python3
"""tools/gen_exam_index.py — the past-paper index (roadmap 5.8, docs/exams-plan.md §4.1).

Reads the user's own evidence folder (default ~/Downloads/A-level, or the
path in KOS_ALEVEL / --root):

  Topic Practice/Question Index.xlsx     Maths, Computer Science, Packs sheets
  Topic Practice/* — PMT/                PMT's topic packs (file names only)
  _build/{maths,cs}/taxonomy.json        the boards' AS / A-level markings

and the hand-checked map tools/exam-index/map.csv (board sub-topic → spec
leaves), and writes js/data/exam-index.js. That file is GENERATED and
committed (invariant 125): it holds references and numbers only — paper,
sitting, question number, marks, leaves, the user's own one-line tag
reason — never past-paper text or images, and file paths relative to the
user's folder for them to open themselves.

It fails loudly, never silently:
  * the folder or the workbook missing        → a clear message, exit 2,
                                                 the committed file untouched
  * a board sub-topic with no row in map.csv  → listed, exit 1
  * a map row naming a leaf that doesn't exist → listed, exit 1

Usage:  python3 tools/gen_exam_index.py [--root PATH] [--check]
        --check  build in memory and compare with the committed file
"""
import csv
import json
import os
import re
import sys
from collections import OrderedDict, defaultdict
from datetime import date

HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.dirname(HERE)
OUT = os.path.join(REPO, "js", "data", "exam-index.js")
MAP = os.path.join(HERE, "exam-index", "map.csv")

MONTHS = {m: i + 1 for i, m in enumerate(["January", "February", "March", "April", "May", "June", "July",
                                         "August", "September", "October", "November", "December"])}
SUBJECT = {"Maths": "maths", "Computer Science": "compsci"}
# (subject, level, paper) → the preset's index id (js/data/exam-presets.js)
PAPER_ID = {
    ("compsci", "A-Level", "Paper 1"): "aqa-7517-1", ("compsci", "A-Level", "Paper 2"): "aqa-7517-2",
    ("compsci", "AS", "Paper 1"): "aqa-7516-1", ("compsci", "AS", "Paper 2"): "aqa-7516-2",
    ("maths", "A-Level", "Paper 1"): "edexcel-9ma0-1", ("maths", "A-Level", "Paper 2"): "edexcel-9ma0-2",
    ("maths", "A-Level", "Paper 3"): "edexcel-9ma0-3",
    ("maths", "AS", "Paper 1"): "edexcel-8ma0-1", ("maths", "AS", "Paper 2"): "edexcel-8ma0-2",
}
PAPER_META = {
    "aqa-7517-1": ("AQA", "7517", "A", "AQA CS Paper 1", 100), "aqa-7517-2": ("AQA", "7517", "A", "AQA CS Paper 2", 100),
    "aqa-7516-1": ("AQA", "7516", "AS", "AQA AS CS Paper 1", 75), "aqa-7516-2": ("AQA", "7516", "AS", "AQA AS CS Paper 2", 75),
    "edexcel-9ma0-1": ("Edexcel", "9MA0", "A", "Edexcel Maths Paper 1", 100), "edexcel-9ma0-2": ("Edexcel", "9MA0", "A", "Edexcel Maths Paper 2", 100),
    "edexcel-9ma0-3": ("Edexcel", "9MA0", "A", "Edexcel Maths Paper 3", 100),
    "edexcel-8ma0-1": ("Edexcel", "8MA0", "AS", "Edexcel AS Maths Paper 1", 100), "edexcel-8ma0-2": ("Edexcel", "8MA0", "AS", "Edexcel AS Maths Paper 2", 60),
}
PMT_MATHS_UNIT = {
    "Proof": "P1", "Algebra-and-Functions": "P2", "Coordinate-Geometry": "P3", "Sequences-and-Series": "P4",
    "Trigonometry": "P5", "Exponentials-and-Logarithms": "P6", "Differentiation": "P7", "Integration": "P8",
    "Numerical-Methods": "P9", "Vectors": "P10", "Statistical-Sampling": "S1", "Data-Presentation-and-Interpretation": "S2",
    "Probability": "S3", "Statistical-Distributions": "S4", "Hypothesis-Testing": "S5",
    "Quantities-and-Units-in-Mechanics": "M6", "Kinematics": "M7", "Forces-and-Newtons-Laws": "M8", "Moments": "M9",
}


def die(msg, code=2):
    sys.stderr.write("gen_exam_index: " + msg + "\n")
    sys.exit(code)


def sitting(session):
    s = str(session or "").strip()
    if s.lower() in ("specimen", "sample"):
        return s.lower()
    m = re.match(r"^([A-Za-z]+)\s+(\d{4})$", s)
    if not m or m.group(1) not in MONTHS:
        return None
    return "%s-%02d" % (m.group(2), MONTHS[m.group(1)])


def sit_key(s):
    return ("0000-00" if s in ("specimen", "sample") else s)


def load_leaves():
    """the generated spec leaves and units, read straight from js/data (no Node needed)"""
    out = {}
    for sid, fn in (("maths", "maths.js"), ("compsci", "compsci.js")):
        src = open(os.path.join(REPO, "js", "data", fn), encoding="utf8").read()
        start = src.index("KOS_DATA." + sid + " =")
        body = src[src.index("{", start):src.rindex("}") + 1]
        data = json.loads(body)
        leaves, units = set(), {}
        def walk(n, depth, unit):
            unit = unit or n["ref"]
            if depth == 0:
                units[n["ref"]] = []
            if n.get("content"):
                leaves.add(n["ref"])
                units[unit].append(n["ref"])
            for c in n.get("children") or []:
                walk(c, depth + 1, unit)
        for sec in data["sections"]:
            walk(sec, 0, None)
        out[sid] = (leaves, units)
    return out


def load_map(leaves):
    rows, bad = {}, []
    with open(MAP, encoding="utf8", newline="") as f:
        for r in csv.DictReader(f):
            sid, sub = r["subject"].strip(), r["subtopic"].strip()
            refs = r["leaves"].split()
            for ref in refs:
                if ref not in leaves[sid][0]:
                    bad.append("%s: %r names %s, which is not a leaf" % (sid, sub, ref))
            rows[(sid, sub)] = refs
    return rows, bad


def main(argv):
    root = os.environ.get("KOS_ALEVEL") or os.path.expanduser("~/Downloads/A-level")
    if "--root" in argv:
        root = argv[argv.index("--root") + 1]
    xlsx = os.path.join(root, "Topic Practice", "Question Index.xlsx")
    if not os.path.isdir(root) or not os.path.isfile(xlsx):
        die("the evidence folder is not on this machine (looked for %s).\n"
            "  js/data/exam-index.js is committed, so the app keeps working; rebuild it on the\n"
            "  machine that has the folder, or pass --root /path/to/A-level." % xlsx)
    try:
        import openpyxl
    except ImportError:
        die("needs openpyxl (pip install --user openpyxl)")

    leaves = load_leaves()
    mapping, bad = load_map(leaves)
    wb = openpyxl.load_workbook(xlsx, read_only=True)

    papers = OrderedDict()
    questions = OrderedDict()           # (paper, sitting, section, q) → row
    unmapped = defaultdict(int)
    for sheet, sid in SUBJECT.items():
        rows = list(wb[sheet].iter_rows(values_only=True))
        head = rows[0]
        ix = {h: i for i, h in enumerate(head)}
        for r in rows[1:]:
            if r[ix["Kind"]] != "QP":
                continue
            pid = PAPER_ID.get((sid, r[ix["Level"]], r[ix["Paper"]]))
            st = sitting(r[ix["Session"]])
            if not pid or not st:
                bad.append("%s row with unknown paper or session: %r" % (sheet, r[:4]))
                continue
            section = r[ix["Section / area"]] or ""
            if sid == "compsci":
                section = "A" if pid in ("aqa-7517-1", "aqa-7516-1") else ""
            q = str(r[ix["Question"]]).strip()
            key = (pid, st, section, q)
            subs = [s for s in (r[ix["Sub-topic 1"]], r[ix["Sub-topic 2"]]) if s]
            refs = []
            for s in subs:
                m = mapping.get((sid, s))
                if m is None:
                    unmapped[(sid, s)] += 1
                    continue
                refs += [sid + ":" + x for x in m if sid + ":" + x not in refs]
            if key not in questions:
                meta = PAPER_META[pid]
                p = papers.setdefault(pid, {"subject": sid, "board": meta[0], "qual": meta[1], "level": meta[2],
                                             "paper": "Paper " + pid[-1], "title": meta[3], "marks": meta[4], "sittings": set()})
                p["sittings"].add(st)
                why = str(r[ix["Why this tag"]] or "").strip()
                questions[key] = {"marks": int(r[ix["Marks"]] or 0), "refs": refs, "subs": subs,
                                  "why": why[:140], "file": r[ix["Source paper"]] or ""}
            else:
                for x in refs:
                    if x not in questions[key]["refs"]:
                        questions[key]["refs"].append(x)
                for s in subs:
                    if s not in questions[key]["subs"]:
                        questions[key]["subs"].append(s)

    # our compiled packs (Packs sheet): QP and MS rows become one pack
    packs = OrderedDict()
    for r in list(wb["Packs"].iter_rows(values_only=True))[1:]:
        subj, area, topic, sub, level, kind, units, pages, file = r[:9]
        sid = SUBJECT.get(subj)
        if not sid:
            continue
        pkid = "ours:%s:%s:%s" % (sid, re.sub(r"[^a-z0-9]+", "-", str(sub).lower()).strip("-"), "as" if level == "AS" else "a")
        m = mapping.get((sid, sub))
        if m is None:
            unmapped[(sid, sub)] += 0
        p = packs.setdefault(pkid, {"id": pkid, "source": "ours", "subject": sid, "area": area or "", "topic": topic or "",
                                    "sub": sub, "level": level, "refs": [sid + ":" + x for x in (m or [])],
                                    "units": 0, "pages": 0, "file": "", "ms": ""})
        if kind == "QP":
            p["units"], p["pages"], p["file"] = int(units or 0), int(pages or 0), file
        else:
            p["ms"] = file

    # PMT's packs: file names only, linked at the unit
    tp = os.path.join(root, "Topic Practice")
    for sid, folder in (("maths", "Maths — PMT"), ("compsci", "Computer Science — PMT")):
        base = os.path.join(tp, folder)
        if not os.path.isdir(base):
            continue
        for dirpath, _, files in sorted(os.walk(base)):
            for fn in sorted(files):
                if not fn.endswith(".pdf") or re.search(r"(Answers|\bMS|Mark ?Scheme|Model|Cheat)", fn, re.I) \
                        or re.search(r"Cheat|Notes|Model", dirpath):
                    continue
                rel = os.path.relpath(os.path.join(dirpath, fn), tp)
                parts = os.path.relpath(dirpath, base).split(os.sep)
                if sid == "maths":
                    m = re.match(r"^(.*)-(\d)$", parts[1] if len(parts) > 1 else "")
                    if not m or m.group(1) not in PMT_MATHS_UNIT:
                        continue
                    unit, level = PMT_MATHS_UNIT[m.group(1)], ("AS" if m.group(2) == "1" else "A-Level")
                    area = parts[0]
                else:
                    if parts[0] != "Topic-Qs" or len(parts) < 2:
                        continue
                    m = re.match(r"^(\d\d)-.*?(-AS)?$", parts[1])
                    if not m:
                        continue
                    unit, level, area = "4.%d" % int(m.group(1)), ("AS" if m.group(2) else "A-Level"), "Topic questions"
                title = re.sub(r"\s+(QP)$", "", os.path.splitext(fn)[0])
                pkid = "pmt:" + re.sub(r"[^a-z0-9]+", "-", rel.lower()).strip("-")[:90]
                packs[pkid] = {"id": pkid, "source": "pmt", "subject": sid, "area": area, "topic": parts[1] if len(parts) > 1 else "",
                               "sub": title, "level": level, "refs": [sid + ":" + unit], "units": 0, "pages": 0,
                               "file": rel, "ms": ""}

    # the A-level-only leaves: every board sub-topic mapping to the leaf is A-level only
    level_of = {}
    for sid, path in (("maths", "maths/taxonomy.json"), ("compsci", "cs/taxonomy.json")):
        tx = json.load(open(os.path.join(root, "_build", path), encoding="utf8"))
        def subs_of(node):
            if isinstance(node, list):
                for s in node:
                    yield s
            elif isinstance(node, dict):
                for k, v in node.items():
                    if not k.startswith("_"):
                        yield from subs_of(v)
        for s in subs_of(tx):
            a2 = (s.get("as") is False) if sid == "maths" else bool(s.get("al"))
            level_of[(sid, s["sub"])] = a2
            if (sid, s["sub"]) not in mapping:
                unmapped[(sid, s["sub"])] += 0
    year2 = {}
    for sid in ("maths", "compsci"):
        verdict = {}
        for (msid, sub), refs in mapping.items():
            if msid != sid or (sid, sub) not in level_of:
                continue
            for ref in refs:
                verdict[ref] = verdict.get(ref, True) and level_of[(sid, sub)]
        year2[sid] = [sid + ":" + r for r in sorted(verdict, key=lambda x: [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", x)]) if verdict[r]]

    if unmapped:
        bad += ["no map.csv row for %s sub-topic %r (%d questions)" % (k[0], k[1], v) for k, v in sorted(unmapped.items())]
    if bad:
        sys.stderr.write("gen_exam_index: the index does not reconcile:\n  " + "\n  ".join(bad) + "\n")
        return 1

    for p in papers.values():
        p["sittings"] = sorted(p["sittings"], key=sit_key)
    qrows = [[k[0], k[1], k[2], k[3], v["marks"], v["refs"], v["why"]] for k, v in questions.items()]
    qrows.sort(key=lambda r: (r[0], sit_key(r[1]), r[2], [int(t) if t.isdigit() else t for t in re.split(r"(\d+)", r[3])]))
    payload = OrderedDict([
        ("v", 1), ("generated", date.today().isoformat()), ("source", "Question Index.xlsx, map.csv"),
        ("papers", papers), ("questions", qrows), ("packs", list(packs.values())), ("year2Only", year2),
    ])
    js = ("/* GENERATED by tools/gen_exam_index.py from the user's own past-paper index —\n"
          "   do not edit by hand (invariant 125). References and numbers only: no\n"
          "   past-paper text or images. `questions` rows are\n"
          "   [paper, sitting, section, question, marks, [leaves], tag reason]. */\n"
          "window.KOS_EXAMS = window.KOS_EXAMS || {};\n"
          "window.KOS_EXAMS.index = " + json.dumps(payload, ensure_ascii=False, separators=(",", ":")) + ";\n")
    js = js.replace('],["', '],\n["')
    if "--check" in argv:
        old = open(OUT, encoding="utf8").read() if os.path.exists(OUT) else ""
        strip = lambda s: re.sub(r'"generated":"[^"]*"', "", s)
        if strip(old) != strip(js):
            sys.stderr.write("gen_exam_index: js/data/exam-index.js is out of date — rerun without --check\n")
            return 1
        print("gen_exam_index: up to date")
        return 0
    open(OUT, "w", encoding="utf8").write(js)
    print("gen_exam_index: %d papers, %d questions, %d packs, year 2 only: maths %d, cs %d → %s" % (
        len(papers), len(qrows), len(packs), len(year2["maths"]), len(year2["compsci"]), os.path.relpath(OUT, REPO)))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
